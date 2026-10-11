import io, json, os, re, secrets, threading, time
from collections import Counter
import requests
from docx import Document
from flask import Flask, render_template, request, send_file, abort, redirect, url_for

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 5 * 1024 * 1024
CACHE, LOCK = {}, threading.Lock()
TTL = 1800
STOP = set('the and with for are you your from that this will have has our their job work years experience strong ability skills knowledge required preferred using within across into about including must should candidate team teams role responsible good excellent develop development support supporting business project projects solutions systems technology technologies professional'.split())
MODES = {'ats':'ATS-Friendly','recruiter':'Recruiter-Friendly','hands-on':'Hands-On Experience','all-in-one':'All-in-One AI Tailored'}

def terms(s):
    return {w for w in re.findall(r'(?<!\w)[a-z][a-z0-9+#./-]{2,}(?!\w)',s.lower()) if w not in STOP}

def paragraphs(doc):
    items = list(doc.paragraphs)
    seen = {id(p._p) for p in items}
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if id(p._p) not in seen:
                        seen.add(id(p._p));items.append(p)
    return items

def extract(doc):
    return '\n'.join(p.text for p in paragraphs(doc))

def stats(jd, text):
    target, actual = terms(jd), terms(text)
    matched, missing = sorted(target & actual),sorted(target-actual)
    return {'score':round(100*len(matched)/max(1,len(target))), 'found':matched[:90], 'missing':missing[:90]}

def bullet(p):
    style=(p.style.name or '').lower()
    return ('list' in style or 'bullet' in style or p.text.lstrip().startswith(('•','-','–','*')) or
            (p._p.pPr is not None and p._p.pPr.numPr is not None))

def relevance(s, wanted, mode):
    score=len(terms(s)&wanted)
    if mode in ('hands-on','all-in-one'):
        score+=sum(v in s.lower() for v in ('implemented','developed','designed','configured','troubleshot','optimized','migrated','automated','integrated','tested','debugged','deployed','built'))
    if mode=='recruiter':score+=min(len(s.split())/30,1.5)
    return score

def reorder(doc,jd,mode):
    ps=list(doc.paragraphs);wanted=terms(jd);i=0
    while i<len(ps):
        if not bullet(ps[i]):i+=1;continue
        j=i
        while j<len(ps) and bullet(ps[j]):j+=1
        group=ps[i:j]
        if len(group)>1:
            parent=group[0]._p.getparent()
            if all(p._p.getparent() is parent for p in group):
                idx=parent.index(group[0]._p)
                for p in group:parent.remove(p._p)
                for off,p in enumerate(sorted(group,key=lambda p:-relevance(p.text,wanted,mode))):parent.insert(idx+off,p._p)
        i=j

def set_text(p, text):
    if p.text==text:return
    if p.runs:
        p.runs[0].text=text
        for run in p.runs[1:]:run.text=''
    else:p.add_run(text)

def create_doc(raw,jd,mode,edits=None):
    doc=Document(io.BytesIO(raw))
    if edits is not None:
        ps=[p for p in paragraphs(doc) if p.text.strip()]
        if len(ps)!=len(edits):raise ValueError('Paragraph count mismatch. Please analyze again.')
        for p,t in zip(ps,edits):set_text(p,t)
    reorder(doc,jd,mode)
    output=io.BytesIO();doc.save(output);return output.getvalue()

def gemini_chunk(rows,jd):
    key=os.environ.get('GEMINI_API_KEY','').strip()
    if not key:raise ValueError('Gemini API key is not configured. Add GEMINI_API_KEY in Render Environment.')
    model=os.environ.get('GEMINI_MODEL','gemini-2.5-flash')
    url=f'https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent'
    prompt='''You are a meticulous resume editor. Tailor to the JD using ONLY verified evidence in the supplied original resume text. Never invent or assume skills, clients, locations, titles, dates, education, numbers, accomplishments or years. Preserve names, headings, client/location/date lines exactly. Improve summaries and bullets into natural recruiter-friendly, specific hands-on phrasing when supported by source. Do not claim 100% ATS match. Keep each paragraph in its original place. Output JSON with key "items" containing objects {"id": original integer id, "text": replacement string}, exactly one for each supplied item. Never combine, split, remove or add paragraphs. If there is insufficient evidence, preserve the original text. Avoid bloating length. Original items are extracted from a Word resume; some are inside tables.\nJOB DESCRIPTION:\n'''+jd[:12000]+'\nRESUME PARAGRAPHS:\n'+json.dumps(rows,ensure_ascii=False)
    try:
        r=requests.post(url,headers={'x-goog-api-key':key},json={'contents':[{'parts':[{'text':prompt}]}],'generationConfig':{'responseMimeType':'application/json','temperature':0.2}},timeout=110)
    except requests.RequestException as e:raise ValueError('Gemini network request failed; try again.') from e
    if not r.ok:raise ValueError(f'Gemini API returned {r.status_code}. Check your API key, model, or free-tier quota.')
    try:
        payload=json.loads(r.json()['candidates'][0]['content']['parts'][0]['text'])
        items=payload['items']; mapping={int(x['id']):str(x['text']) for x in items}
        if set(mapping)!={x['id'] for x in rows}:raise ValueError('Incomplete AI response')
        return mapping
    except (KeyError,ValueError,IndexError,TypeError) as e:raise ValueError('Gemini response was incomplete; please retry.') from e

def ai_rewrite(raw,jd):
    ps=[p for p in paragraphs(Document(io.BytesIO(raw))) if p.text.strip()]
    if len(ps)>260:raise ValueError('Resume has too many paragraphs for this free version (maximum 260).')
    originals=[p.text for p in ps];results=list(originals)
    # Chunk to reduce output truncation; free API quota may still limit calls.
    for start in range(0,len(ps),14):
        batch=[{'id':i,'text':originals[i]} for i in range(start,min(start+14,len(ps)))]
        updated=gemini_chunk(batch,jd)
        for i,text in updated.items():results[i]=text
    return results

def get_entry(token):
    with LOCK:
        e=CACHE.get(token)
        if e and time.time()-e['created']<TTL:return e
    abort(404,description='Resume session expired. Upload again.')

def remember(raw,jd):
    token=secrets.token_urlsafe(32)
    with LOCK:
        now=time.time()
        for k in list(CACHE):
            if now-CACHE[k]['created']>TTL:del CACHE[k]
        CACHE[token]={'created':now,'raw':raw,'jd':jd,'edits':None,'suggestions':None,'ai_mode':'auto'}
    return token

@app.get('/')
def home():return render_template('index.html')

@app.post('/analyze')
def analyze():
    jd=request.form.get('jd','').strip();f=request.files.get('resume')
    if not jd or not f or not (f.filename or '').lower().endswith('.docx'):
        return render_template('index.html',error='Upload a .docx resume and paste a JD.',jd=jd),400
    raw=f.read()
    try:
        doc=Document(io.BytesIO(raw));text=extract(doc)
        if not text.strip():raise ValueError()
    except Exception:return render_template('index.html',error='Unable to read this DOCX resume.',jd=jd),400
    token=remember(raw,jd)
    return render_template('index.html',token=token,jd=jd,preview=text[:3500],message='Estimated keyword overlap, not an official ATS score.',**stats(jd,text))

@app.get('/download/<token>/<mode>')
def download(token,mode):
    if mode not in MODES:abort(404)
    e=get_entry(token)
    if mode=='all-in-one' and e['edits'] is None:return redirect(url_for('review',token=token))
    output=create_doc(e['raw'],e['jd'],mode,e['edits'] if mode=='all-in-one' else None)
    return send_file(io.BytesIO(output),as_attachment=True,download_name=f'CareerLift_{mode}.docx',mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@app.post('/generate/<token>')
def generate(token):
    e=get_entry(token)
    ai_mode=request.form.get('ai_mode','auto')
    if ai_mode not in ('manual','auto'):abort(400)
    try:edits=ai_rewrite(e['raw'],e['jd'])
    except ValueError as exc:
        doc=Document(io.BytesIO(e['raw']));text=extract(doc)
        return render_template('index.html',token=token,jd=e['jd'],preview=text[:3500],error=str(exc),message='Estimated keyword overlap.',**stats(e['jd'],text)),502
    originals=[p.text for p in paragraphs(Document(io.BytesIO(e['raw']))) if p.text.strip()]
    with LOCK:
        e['suggestions']=edits
        e['ai_mode']=ai_mode
        e['edits']=edits if ai_mode=='auto' else list(originals)
    if ai_mode=='auto':
        output=create_doc(e['raw'],e['jd'],'all-in-one',edits)
        return send_file(io.BytesIO(output),as_attachment=True,download_name='Updated_Resume.docx',mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    return redirect(url_for('review',token=token))

@app.get('/review/<token>')
def review(token):
    e=get_entry(token)
    if e['edits'] is None:return redirect(url_for('home'))
    original=extract(Document(io.BytesIO(e['raw'])))
    tailored='\n'.join(e['edits'])
    return render_template('review.html',token=token,edits=e['edits'],suggestions=e.get('suggestions'),ai_mode=e.get('ai_mode','auto'),before=stats(e['jd'],original),after=stats(e['jd'],tailored))

@app.post('/save/<token>')
def save(token):
    e=get_entry(token)
    if e['edits'] is None:abort(404)
    edits=request.form.getlist('paragraph')
    if len(edits)!=len(e['edits']):abort(400)
    with LOCK:
        e['edits']=edits
    return redirect(url_for('review',token=token))

if __name__=='__main__':app.run(host='0.0.0.0',port=int(os.environ.get('PORT',5000)),debug=False)

@app.post('/apply-suggestions/<token>')
def apply_suggestions(token):
    e=get_entry(token)
    if e.get('ai_mode')!='manual' or not e.get('suggestions'):abort(400)
    originals=[p.text for p in paragraphs(Document(io.BytesIO(e['raw']))) if p.text.strip()]
    selected=set(request.form.getlist('selected'))
    with LOCK:
        e['edits']=[e['suggestions'][i] if str(i) in selected else originals[i] for i in range(len(originals))]
    return redirect(url_for('review',token=token))
